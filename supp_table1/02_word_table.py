"""Write Supplementary Table 1 (stats_table.csv from 01_stats_table.R) as a landscape Word table."""
import os
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Supplementary Table 1 - statistics.docx")
SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")
FIG_ORDER = {"Fig. 4": 0, "Fig. 5": 1, "Fig. 6": 2, "Suppl. Fig. 3": 3, "Suppl. Fig. 4": 4}
CMP_ORDER = {"WT vs GtdKO+AAV-Gαt2 (pre)": 0, "WT vs GtdKO+AAV-Gαt2 (post)": 1, "GtdKO+AAV-Gαt2 pre vs post": 2,
             "WT pre vs post": 3, "Group × training (change, WT vs GtdKO+AAV-Gαt2)": 4}
GROUPS = {"WT vs GtdKO+AAV-Gαt2 (pre)": ("WT", "GtdKO+AAV-Gαt2 pre"),
          "WT vs GtdKO+AAV-Gαt2 (post)": ("WT", "GtdKO+AAV-Gαt2 post"),
          "GtdKO+AAV-Gαt2 pre vs post": ("GtdKO+AAV-Gαt2 pre", "GtdKO+AAV-Gαt2 post"),
          "WT pre vs post": ("WT pre", "WT post"),
          "Group × training (change, WT vs GtdKO+AAV-Gαt2)": ("WT change", "GtdKO+AAV-Gαt2 change")}
TERM = {"group": "Group", "training": "Training", "group:training": "Group × training"}


def p_text(p):
    if p >= 0.001:
        return f"{p:.3f}" if p >= 0.01 else f"{p:.4f}"
    m, e = f"{p:.2e}".split("e")
    return f"{m} × 10{str(int(e)).translate(SUP)}"


def df_text(x):
    return f"{x:.0f}" if abs(x - round(x)) < 1e-6 else f"{x:.1f}"


def shade(cell, hex_):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_)
    tc.append(s)


def runfont(run, size, bold=False, italic=False):
    run.font.name = "Times New Roman"; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")


def para(doc, text, size=10, bold=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            runfont(p.add_run(chunk), size, bold=bold or i % 2 == 1, italic=italic)
    return p


t = pd.read_csv(os.path.join(HERE, "stats_table.csv"))
t["fo"] = t.figure.map(FIG_ORDER); t["co"] = t.comparison.map(CMP_ORDER)
t = t.sort_values(["fo", "panel", "co"], kind="stable")

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.5))

para(doc, "**Supplementary Table 1. Statistics for the neural activity analyses in Figs. 4–6 and Supplementary Figs. 3 "
          "and 4.**", size=11, space_after=6)
para(doc, "All tests were performed on animal-level values (one imaging session per mouse and time point) with "
          "aligned-rank-transform (ART) mixed models (ARTool package in R), with mouse as a random intercept. "
          "Comparisons between groups used value ~ group × level + (1 | mouse); comparisons before and after training "
          "used value ~ training × level + (1 | mouse); the difference in training effect between groups used value ~ "
          "group × training × level + (1 | mouse). Level is the x-axis variable of the panel (spatial frequency, "
          "orientation, or number of neurons for the decoding curves). The table reports the F test for the term "
          "listed. Partial η² = F·df1 / (F·df1 + df2). Means ± SEM are across mice, after averaging each mouse over "
          "the levels of the panel; for the group × training rows, they are the change from before to after training "
          "(post − pre). Neurons are the totals over all sessions in each group (number of visually responsive "
          "neurons is not shown). WT in Figs. 4–6 and Supplementary Fig. 4 is the WT group before training. Units: "
          "responsive fractions, fraction of neurons; OSI, dimensionless; decoding, accuracy (fraction correct); "
          "distances, as in the corresponding panel.", size=9, space_after=8)

cols = ["Figure", "Panel", "Measure", "Comparison", "Term", "n mice\n(1 / 2)", "Neurons\n(1 / 2)",
        "Group 1\nmean ± SEM", "Group 2\nmean ± SEM", "F (df1, df2)", "p", "Partial η²"]
widths = [2.0, 1.2, 3.4, 3.5, 1.8, 1.2, 1.9, 2.3, 2.3, 2.1, 2.0, 1.3]
tab = doc.add_table(rows=1, cols=len(cols))
tab.style = "Table Grid"; tab.alignment = WD_TABLE_ALIGNMENT.CENTER; tab.autofit = False
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tab._tbl.tblPr.append(lay)
for c, wd in zip(tab.columns, widths):
    c.width = Cm(wd)
for c, (h, wd) in enumerate(zip(cols, widths)):
    cell = tab.rows[0].cells[c]; cell.width = Cm(wd); shade(cell, "D9D9D9")
    cell.paragraphs[0].text = ""
    runfont(cell.paragraphs[0].add_run(h), 8, bold=True)
trpr = tab.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trpr.append(hdr)

prev = None
for r in t.itertuples():
    g1, g2 = GROUPS[r.comparison]
    vals = [r.figure, r.panel, r.measure, f"{g1} vs {g2}", TERM[r.term], r.n_mice, r.neurons.replace("+", " + "),
            r.mean_sem_1, r.mean_sem_2, f"{r.F:.2f} ({df_text(r.df1)}, {df_text(r.df2)})", p_text(r.p),
            f"{r.eta2p:.3f}"]
    row = tab.add_row()
    cs = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cs)
    for c, (v, wd) in enumerate(zip(vals, widths)):
        cell = row.cells[c]; cell.width = Cm(wd)
        pp = cell.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
        runfont(pp.add_run(str(v)), 8, bold=(c == 10 and r.p < 0.05))
    if prev is not None and r.figure != prev:
        for cell in row.cells:
            b = OxmlElement("w:tcBorders"); top = OxmlElement("w:top")
            top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "12"); top.set(qn("w:color"), "000000")
            b.append(top); cell._tc.get_or_add_tcPr().append(b)
    prev = r.figure
para(doc, "", size=4)
para(doc, "p < 0.05 in bold. GtdKO+AAV-Gαt2, GtdKO mice with AAV-mediated Gαt2 rescue; SF, spatial frequency; OSI, "
          "orientation selectivity index (cross-validated, see Methods).", size=9)
doc.save(OUT)
print("saved", os.path.abspath(OUT), len(t), "rows")
